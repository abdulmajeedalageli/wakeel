"""Listings from the owner's Excel workbook, loaded via Telegram, and the privacy rules:
customers are never told a price, a person's name or a phone number."""

import asyncio
import io
import json

import pytest
from openpyxl import Workbook

from app import owner_console
from app.agent import Agent, clean_reply
from app.config import get_settings
from app.excel_listings import ExcelConfig, ExcelFormatError, read_listings
from app.locations import Locations
from app.properties import EXCEL_CONFIG_FILE, LOCATIONS_FILE, SHEET_CONFIG_FILE, PropertyStore, listings_file_for
from app.sheet import SheetConfig
from tests.conftest import OWNER_TELEGRAM_ID
from tests.test_agent import FakeGemini, msg, text_response, tool_call_response

SALE_HEADERS = ["رقم العقار", "النوع", "الحي", "الواجهة", "المساحة", "السعر", "المالك", "رقم الوسيط", "حصري",
                "رقم المالك", "عمر العقار", "حالة السوق", "المسؤول", "رقم اللوحة", "الملاحظات"]
RENT_HEADERS = ["رقم العقار", "نوع العقار", "الحي", "الإيجار السنوي", "المالك", "رقم الجوال", "ملاحظات",
                "حالة التأجير", "المسؤول", "الملاحظات", "حالة الطلب"]
# Private values planted in the fake workbook: none may ever reach the AI.
SECRETS = ["مالك سري", "0549998888", "0551112222", "موظف سري", "ملاحظة سرية", "عميل سري", "مصروف سري", "7777",
           "290 ألف", "290000", "11الف", "11000", "4800", "لوحة سرية"]


def build_workbook(extra_tabs: bool = True) -> bytes:
    wb = Workbook()
    sale = wb.active
    sale.title = "عروض البيع (سكني)"
    sale.append(SALE_HEADERS)
    sale.append([1, "فلة", "الذهبي", "", "315م", "", "مالك سري", "0549998888", "غير حصري", "0549998888",
                 "10 سنوات", "تم البيع", "موظف سري", "", "ملاحظة سرية"])
    sale.append([2, "أرض سكنية", "النهضة", "شمالية شارع 20", "400م", "290 ألف", "مالك سري", "", "حصري", "",
                 "", "شاغر", "موظف سري", "لوحة سرية", "ملاحظة سرية"])
    sale.append([3, "استراحة", "الصحنة - الشارع العام", "غربية", "3الف", "", "", "", "", "", "", "", "", "", ""])
    sale.append([None, None, None, None, None, None, None, None, None, None, None, None, None, None, None])
    commercial = wb.create_sheet("عروض البيع (تجاري + زراعي)")  # spacing differs from config: still found
    commercial.append(SALE_HEADERS)
    commercial.append([1, "مزرعه", "الخرج", "", "", "", "", "", "", "", "", "", "", "", "ملاحظة سرية"])
    rent = wb.create_sheet("عروض التأجير (سكني)")
    rent.append(RENT_HEADERS)
    rent.append([1, "شقة", "شارع الثلاثين", "11الف", "مالك سري", "0551112222", "ملاحظة سرية", "سنوي", "موظف سري",
                 "ملاحظة سرية", ""])
    rent.append([2, "غرف عمال", "الصناعية", "400 شهري", "", "", "", "شهري", "", "", "منتهي"])
    rent_c = wb.create_sheet("عروض التأجير (تجاري)")
    rent_c.append(RENT_HEADERS)
    rent_c.append([1, "محل تجاري", "الصناعية", "", "", "", "", "", "", "", ""])
    if extra_tabs:
        requests = wb.create_sheet("طلبات الشراء (سكني)")
        requests.append(["رقم الطلب", "اسم العميل", "الجوال", "نوع الطلب", "المساحة", "الحي", "الميزانية"])
        requests.append([1, "عميل سري", "0551112222", "أرض سكنية", "", "النهضة", "مليون"])
        finance = wb.create_sheet("المالية")
        finance.append(["المصروفات", "المصروفات"])
        finance.append(["مصروف سري", 7777])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.fixture
def excel_store(tmp_path) -> PropertyStore:
    store = PropertyStore(None, SheetConfig.load(SHEET_CONFIG_FILE), Locations.load(LOCATIONS_FILE))
    store.replace_with_excel(build_workbook(), tmp_path / "listings.xlsx")
    return store


def run(coro):
    return asyncio.run(coro)


# --- reading the workbook ---

def test_only_the_listing_tabs_are_read(excel_store):
    tabs = {p.source_tab for p in excel_store.properties}
    assert tabs == {"عروض البيع (سكني)", "عروض البيع (تجاري+ زراعي)", "عروض التأجير (سكني)", "عروض التأجير (تجاري)"}
    everything = json.dumps([p.to_dict() for p in excel_store.properties], ensure_ascii=False)
    # Not even loaded into memory: owner names/phones, staff, notes, requests, finances, plot numbers.
    for secret in ("مالك سري", "0549998888", "0551112222", "موظف سري", "ملاحظة سرية", "عميل سري", "مصروف سري",
                   "لوحة سرية"):
        assert secret not in everything, secret


def test_ids_types_status_and_numbers(excel_store):
    by_id = {p.property_id: p for p in excel_store.properties}
    assert set(by_id) == {"بيع سكني 1", "بيع سكني 2", "بيع سكني 3", "بيع تجاري 1", "إيجار سكني 1", "إيجار سكني 2",
                          "إيجار تجاري 1"}
    villa, land, rest = by_id["بيع سكني 1"], by_id["بيع سكني 2"], by_id["بيع سكني 3"]
    assert (villa.type, villa.status, villa.age) == ("villa", "sold", "10 سنوات")
    assert (land.type, land.status, land.price_sar, land.area_sqm, land.facade) == (
        "land", "available", 290000, 400, "شمالية شارع 20")
    assert (rest.type, rest.status, rest.area_sqm, rest.price_sar) == ("rest_house", "available", 3000, None)
    assert by_id["بيع تجاري 1"].type == "farm"
    flat, labor = by_id["إيجار سكني 1"], by_id["إيجار سكني 2"]
    assert (flat.deal, flat.price_sar, flat.terms, flat.status) == ("rent", 11000, "سنوي", "available")
    assert (labor.type, labor.price_sar, labor.status) == ("labor_housing", 4800, "rented")  # 400 monthly x 12
    assert by_id["إيجار تجاري 1"].status == "available"  # empty status = available


def test_ids_find_the_right_listing_across_tabs(excel_store):
    assert excel_store.get("بيع سكني 1").source_tab == "عروض البيع (سكني)"
    assert excel_store.get("إيجار سكني 1").source_tab == "عروض التأجير (سكني)"
    assert excel_store.get("بيع سكني ١").property_id == "بيع سكني 1"  # Arabic digits


def test_not_an_excel_file_or_no_listing_tabs():
    config, locations = ExcelConfig.load(EXCEL_CONFIG_FILE), Locations.load(LOCATIONS_FILE)
    with pytest.raises(ExcelFormatError):
        read_listings(b"this is not excel", config, locations)
    wb = Workbook()
    wb.active.title = "ورقة1"
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(ExcelFormatError):
        read_listings(out.getvalue(), config, locations)


def test_a_bad_file_keeps_the_old_listings(excel_store, tmp_path):
    before = len(excel_store.properties)
    with pytest.raises(ExcelFormatError):
        excel_store.replace_with_excel(b"broken", tmp_path / "listings.xlsx")
    assert len(excel_store.properties) == before
    assert (tmp_path / "listings.xlsx").read_bytes()[:2] == b"PK"  # the good file wasn't overwritten


# --- searching ---

def test_district_inside_a_descriptive_cell_is_found(excel_store):
    result = excel_store.search(district="الصحنة")
    assert [p.property_id for p in result.properties] == ["بيع سكني 3"]


def test_budget_filters_but_unknown_prices_stay_in(excel_store):
    ids = {p.property_id for p in excel_store.search(deal="sale", max_price=200000).properties}
    assert "بيع سكني 2" not in ids  # 290,000 is over budget
    assert {"بيع سكني 3", "بيع تجاري 1"} <= ids  # no price: the team handles it


def test_the_uploaded_file_is_used_after_a_restart(tmp_path, monkeypatch):
    from app.properties import build_property_store

    (tmp_path / "listings.xlsx").write_bytes(build_workbook())
    settings = get_settings().model_copy(update={"db_path": str(tmp_path / "bot.db"), "google_sheet_id": None})
    store = build_property_store(settings)
    assert store.source == "excel" and store.refresh()
    assert len(store.properties) == 7


# --- privacy: what the AI sees ---

def test_no_price_name_phone_or_note_ever_reaches_the_ai(excel_store):
    fake = FakeGemini(
        tool_call_response(("search_properties", {"deal": "sale"}), ("search_properties", {"deal": "rent"})),
        tool_call_response(*[("get_property_details", {"property_id": p.property_id})
                             for p in excel_store.properties]),
        text_response("عندنا أرض في النهضة"),
    )
    reply = run(Agent(get_settings(), excel_store, client=fake).reply([msg("in", "كم سعر الأرض ومين المالك؟")]))
    assert reply.ok
    for i in range(len(fake.requests)):
        sent = fake.request_json(i)
        for secret in SECRETS:
            assert secret not in sent, f"request {i} leaked {secret!r}"
    assert "never state" in fake.request_json(1)  # the price rule travels with the results


def test_phone_numbers_never_leave_in_a_reply():
    assert clean_reply("تواصل مع المالك على 0549998888 أو +966 55 111 2222") == "تواصل مع المالك على أو"
    assert clean_reply("المساحة 315م ورقم العقار بيع سكني 2") == "المساحة 315م ورقم العقار بيع سكني 2"


def test_price_questions_have_their_own_alert_reason():
    assert owner_console.ALERT_REASONS["price_request"] == "يسأل عن السعر"


# --- loading the file from Telegram ---

GROUP = -1001234567890


def excel_message(sender=OWNER_TELEGRAM_ID, thread_id=None, file_id="X1", name="عروض فريم.xlsx") -> dict:
    message = {"message_id": 9, "from": {"id": sender}, "chat": {"id": GROUP},
               "document": {"file_id": file_id, "file_name": name,
                            "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}}
    if thread_id is not None:
        message.update({"message_thread_id": thread_id, "is_topic_message": True})
    return {"update_id": 1, "message": message}


@pytest.fixture
def group(make_state, fake_telegram, tmp_path):
    fake_telegram.chat_id = GROUP
    fake_telegram.topics_enabled = True
    state = make_state(db_path=str(tmp_path / "bot.db"))
    state.properties = PropertyStore(None, SheetConfig.load(SHEET_CONFIG_FILE), Locations.load(LOCATIONS_FILE))
    return state


def test_sending_the_excel_in_general_loads_the_listings(group, fake_telegram, tmp_path):
    fake_telegram.downloadable["X1"] = build_workbook()
    run(owner_console.handle_update(excel_message(), group))
    text = fake_telegram.sent[-1][0]
    assert text.startswith("✅ تم تحميل العقارات: 7 (5 متاح)")
    assert "• عروض البيع (سكني): 3" in text
    assert group.properties.source == "excel" and len(group.properties.properties) == 7
    assert listings_file_for(str(tmp_path / "bot.db")).exists()  # kept for the next restart
    run(owner_console.handle_update({"update_id": 2, "message": {
        "message_id": 10, "from": {"id": OWNER_TELEGRAM_ID}, "chat": {"id": GROUP}, "text": "/status"}}, group))
    assert "🏠 العقارات: 7 (5 متاح) من ملف الإكسل" in fake_telegram.sent[-1][0]


def test_a_broken_file_is_refused_and_the_old_listings_stay(group, fake_telegram):
    fake_telegram.downloadable["X1"] = build_workbook()
    run(owner_console.handle_update(excel_message(), group))
    fake_telegram.downloadable["BAD"] = b"not excel at all"
    run(owner_console.handle_update(excel_message(file_id="BAD"), group))
    assert fake_telegram.sent[-1][0].startswith("❌ ما قدرت أقرأ الملف")
    assert len(group.properties.properties) == 7


def test_an_excel_in_a_customers_topic_goes_to_the_customer(group, fake_telegram, fake_wa, db):
    db.get_or_create_chat("966512341234")
    thread = run(owner_console.customer_thread(group, "966512341234"))
    fake_telegram.downloadable["X1"] = build_workbook()
    run(owner_console.handle_update(excel_message(thread_id=thread), group))
    assert fake_wa.sent_media and fake_wa.sent_media[0][1] == "document"  # relayed, not loaded
    assert group.properties.properties == []


def test_someone_not_on_the_team_cannot_replace_the_listings(group, fake_telegram):
    fake_telegram.downloadable["X1"] = build_workbook()
    run(owner_console.handle_update(excel_message(sender=999), group))
    assert group.properties.properties == [] and fake_telegram.sent == []
